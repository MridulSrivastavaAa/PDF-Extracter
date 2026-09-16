"""
Styled Excel Exporter for MoSPI Project Data.
Formats the canonical 20-column dataframe into professional Excel spreadsheets
with Navy Blue headers (#1F4E79), zebra striping (#F2F5F9), thin borders, and freeze panes.
"""

import os
from typing import List, Dict, Any
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from .schema import CANONICAL_COLUMNS

HEADER_FILL = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
ZEBRA_FILL = PatternFill(start_color="F2F5F9", end_color="F2F5F9", fill_type="solid")
WHITE_FILL = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
THIN_BORDER_SIDE = Side(border_style="thin", color="D9D9D9")
ROW_BORDER = Border(left=THIN_BORDER_SIDE, right=THIN_BORDER_SIDE, top=THIN_BORDER_SIDE, bottom=THIN_BORDER_SIDE)

def export_records_to_styled_excel(records: List[Dict[str, Any]], output_path: str, sheet_title: str = "Extracted_Project_Data") -> str:
    """
    Saves records into a beautifully styled Excel workbook matching the master schema.
    """
    os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
    
    # Create DataFrame and ensure canonical column ordering
    df = pd.DataFrame(records)
    for col in CANONICAL_COLUMNS:
        if col not in df.columns:
            df[col] = None
    df = df[CANONICAL_COLUMNS]
    
    # Write to Excel
    with pd.ExcelWriter(output_path, engine='openpyxl') as writer:
        df.to_excel(writer, index=False, sheet_name=sheet_title)
        
    # Open and apply styling
    wb = openpyxl.load_workbook(output_path)
    ws = wb[sheet_title]
    
    ws.views.sheetView[0].showGridLines = True
    ws.freeze_panes = 'C2'
    
    # Style Header Row
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = ROW_BORDER
    ws.row_dimensions[1].height = 28
    
    col_widths = {}
    for row_idx, row in enumerate(ws.iter_rows(min_row=2, max_row=ws.max_row), start=2):
        is_even = (row_idx % 2 == 0)
        current_fill = ZEBRA_FILL if is_even else WHITE_FILL
        for col_idx, cell in enumerate(row, start=1):
            cell.fill = current_fill
            cell.border = ROW_BORDER
            c_name = ws.cell(row=1, column=col_idx).value or ''
            val_str = str(cell.value) if cell.value is not None else ''
            
            # Alignments
            if any(k in str(c_name).lower() for k in ['cost', 'expenditure', 'progress']):
                cell.alignment = Alignment(horizontal="right", vertical="center")
            elif any(k in str(c_name).lower() for k in ['id', 'pmgid', 'date', 'month', 'year', 'flag', 'milestone', 'order']):
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")
                
            cur_max = col_widths.get(col_idx, 10)
            col_widths[col_idx] = max(cur_max, min(len(val_str) + 3, 50))
            
    for col_idx, width in col_widths.items():
        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = max(width, 12)
        
    wb.save(output_path)
    wb.close()
    return output_path
